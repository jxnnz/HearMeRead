from typing import Optional, List

from fastapi import APIRouter, Depends, Query, UploadFile, File, Form, status, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel
from sqlalchemy.ext.asyncio import AsyncSession

from app.dependencies import get_current_teacher
from app.db import get_db
from app.models import Teacher, Language, GradeLevel
from app.schema import (
    PassageCreate,
    PassageUpdate,
    PassageResponse,
    PassageListResponse,
    QuestionResponse,
)
from app.services import passage_service, question_service, storage_service
from app.utils.docx_parser import validate_upload, parse_passage_only


router = APIRouter(prefix="/passages", tags=["Passages"])


# ── Combined downloadable templates (one per grade/language variant) ──────────
# Format matches what frontend fileParser.js expects.

_COMBINED_GRADE1_FILIPINO = """\
================================================================================
IMPORTANT NOTICE / PAUNAWA:
1. Only fill out the sections you need (Assessment 1, Assessment 2, or both).
2. Leave unused sections completely blank.
3. If only Assessment 1 is filled out, only Assessment 1 will be extracted.
4. If only Assessment 2 is filled out, only Assessment 2 will be extracted.
5. If BOTH Assessment 1 and Assessment 2 are filled out, both will be extracted!
================================================================================

------------------- SAMPLE FORMAT / MGA HALIMBAWA (DO NOT EDIT SAMPLE) -------------------
Language: Filipino
Grade: 1

--- SAMPLE ASSESSMENT 1 ---
Task 1:
b, ng, T, e, p, s, H, G, u, L
Task 2:
W: sanay, tunay | R: Oo
W: ulam, anim | R: Hindi
Task 2 Sentences:
Ang bata ay masaya. Siya ay mabait. Mahal niya ang kanyang pamilya.

--- SAMPLE ASSESSMENT 2 ---
Story Number: 1
Title: Ang Pagong at ang Matsing
Content: Isulat dito ang buong teksto ng kwento.
Questions:
Q: Sino ang pangunahing tauhan ng kwento?
A: Ang Pagong at ang Matsing
--------------------------------------------------------------------------------

=================== FILL OUT YOUR CONTENT BELOW (ISULAT ANG NILALAMAN DITO) ===================

Language:
Filipino

Grade:
1

--- FILLABLE ASSESSMENT 1 ---
Task 1:


Task 2:
W:
R:

W:
R:


--- FILLABLE ASSESSMENT 2 (STORY & QUESTIONS) ---
Story Number:


Title:


Content:


Questions:
Q:
A:
"""

_COMBINED_GRADE2_FILIPINO = """\
================================================================================
IMPORTANT NOTICE / PAUNAWA:
1. Only fill out the sections you need (Assessment 1, Assessment 2, or both).
2. Leave unused sections completely blank.
3. If only Assessment 1 is filled out, only Assessment 1 will be extracted.
4. If only Assessment 2 is filled out, only Assessment 2 will be extracted.
5. If BOTH Assessment 1 and Assessment 2 are filled out, both will be extracted!
================================================================================

------------------- SAMPLE FORMAT / MGA HALIMBAWA (DO NOT EDIT SAMPLE) -------------------
Language: Filipino
Grade: 2

--- SAMPLE ASSESSMENT 1 ---
Task 1:
aso, bata, kuya, isda, damit, bahay, paaralan, mahal, tahimik, maganda
Task 2 Words:
aklat, lapis, mesa, silya, kotse, puno, bundok, ilog, dagat, langit
Task 2 Sentences:
Ang bata ay pumunta sa paaralan. Siya ay nagdala ng kanyang bag. Masaya siya sa klase.

--- SAMPLE ASSESSMENT 2 ---
Story Number: 1
Title: Ang Pagong at ang Matsing
Content: Isulat dito ang buong teksto ng kwento.
Questions:
Q: Sino ang pangunahing tauhan ng kwento?
A: Ang Pagong at ang Matsing
--------------------------------------------------------------------------------

=================== FILL OUT YOUR CONTENT BELOW (ISULAT ANG NILALAMAN DITO) ===================

Language:
Filipino

Grade:
2

--- FILLABLE ASSESSMENT 1 ---
Task 1:


Task 2 Words:


Task 2 Sentences:


--- FILLABLE ASSESSMENT 2 (STORY & QUESTIONS) ---
Story Number:


Title:


Content:


Questions:
Q:
A:
"""

_COMBINED_GRADE3_FILIPINO = """\
================================================================================
IMPORTANT NOTICE / PAUNAWA:
1. Only fill out the sections you need (Assessment 1, Assessment 2, or both).
2. Leave unused sections completely blank.
3. If only Assessment 1 is filled out, only Assessment 1 will be extracted.
4. If only Assessment 2 is filled out, only Assessment 2 will be extracted.
5. If BOTH Assessment 1 and Assessment 2 are filled out, both will be extracted!
================================================================================

------------------- SAMPLE FORMAT / MGA HALIMBAWA (DO NOT EDIT SAMPLE) -------------------
Language: Filipino
Grade: 3

--- SAMPLE ASSESSMENT 1 ---
Task 1:
magulang, kaibigan, kalikasan, pamayanan, kasipagan, katapatan, pagmamahal, pagiging, katahimikan, responsibilidad
Task 2 Words:
naglalaro, kumakain, nagaaral, tumatakbo, nagtatrabaho, natutulog, nagbabasa, sumusulat, naglalakad, nagtatanong
Task 2 Sentences:
Ang mga bata ay masayang naglalaro sa parke tuwing hapon. Tinutulungan nila ang isa't isa sa oras ng pangangailangan.

--- SAMPLE ASSESSMENT 2 ---
Story Number: 1
Title: Ang Pagong at ang Matsing
Content: Isulat dito ang buong teksto ng kwento.
Questions:
Q: Sino ang pangunahing tauhan ng kwento?
A: Ang Pagong at ang Matsing
--------------------------------------------------------------------------------

=================== FILL OUT YOUR CONTENT BELOW (ISULAT ANG NILALAMAN DITO) ===================

Language:
Filipino

Grade:
3

--- FILLABLE ASSESSMENT 1 ---
Task 1:


Task 2 Words:


Task 2 Sentences:


--- FILLABLE ASSESSMENT 2 (STORY & QUESTIONS) ---
Story Number:


Title:


Content:


Questions:
Q:
A:
"""

_COMBINED_GRADE3_ENGLISH = """\
================================================================================
IMPORTANT NOTICE / PAUNAWA:
1. Only fill out the sections you need (Assessment 1, Assessment 2, or both).
2. Leave unused sections completely blank.
3. If only Assessment 1 is filled out, only Assessment 1 will be extracted.
4. If only Assessment 2 is filled out, only Assessment 2 will be extracted.
5. If BOTH Assessment 1 and Assessment 2 are filled out, both will be extracted!
================================================================================

------------------- SAMPLE FORMAT / MGA HALIMBAWA (DO NOT EDIT SAMPLE) -------------------
Language: English
Grade: 3

--- SAMPLE ASSESSMENT 1 ---
Task 1:
beautiful, environment, community, responsibility, friendship, knowledge, adventure, imagination, celebration, determination
Task 2 Words:
running, jumping, playing, reading, writing, eating, sleeping, helping, listening, learning

--- SAMPLE ASSESSMENT 2 ---
Story Number: 1
Title: The Clever Turtle
Content: Write the full story text here.
Questions:
Q: Who is the main character?
A: The Turtle
--------------------------------------------------------------------------------

=================== FILL OUT YOUR CONTENT BELOW (ISULAT ANG NILALAMAN DITO) ===================

Language:
English

Grade:
3

--- FILLABLE ASSESSMENT 1 ---
Task 1:


Task 2 Words:


--- FILLABLE ASSESSMENT 2 (STORY & QUESTIONS) ---
Story Number:


Title:


Content:


Questions:
Q:
A:
"""

# Map of (grade, language) → combined template string
_COMBINED_TEMPLATES = {
    ("grade_1", "filipino"): _COMBINED_GRADE1_FILIPINO,
    ("grade_2", "filipino"): _COMBINED_GRADE2_FILIPINO,
    ("grade_3", "filipino"): _COMBINED_GRADE3_FILIPINO,
    ("grade_3", "english"):  _COMBINED_GRADE3_ENGLISH,
}


# ── List ──────────────────────────────────────────────────────────────────────
@router.get("", response_model=PassageListResponse, summary="List passages")
async def list_passages(
    page:             int                    = Query(1, ge=1),
    page_size:        int                    = Query(20, ge=1, le=100),
    language:         Optional[Language]     = Query(None),
    grade_level:      Optional[GradeLevel]   = Query(None),
    assessment_type:  Optional[int]          = Query(None, ge=1, le=2),
    include_archived: bool                   = Query(False),
    db:               AsyncSession           = Depends(get_db),
    current_teacher:  Teacher                = Depends(get_current_teacher),
):
    total, passages = await passage_service.get_passages(
        db=db,
        teacher_id=current_teacher.id,
        page=page,
        page_size=page_size,
        language=language,
        grade_level=grade_level,
        include_archived=include_archived,
        assessment_type=assessment_type,
    )
    return PassageListResponse(total=total, page=page, page_size=page_size, passages=passages)


# ── Create manually ───────────────────────────────────────────────────────────
@router.post(
    "",
    response_model=PassageResponse,
    status_code=status.HTTP_201_CREATED,
    summary="Create a passage manually",
)
async def create_passage(
    data:            PassageCreate,
    db:              AsyncSession = Depends(get_db),
    current_teacher: Teacher      = Depends(get_current_teacher),
):
    return await passage_service.create_passage(db=db, data=data, teacher_id=current_teacher.id)


def _set_docx_cell_bg(cell, hex_color: str):
    """Set cell background color in python-docx table."""
    try:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), hex_color)
        tcPr.append(shd)
    except Exception:
        pass


def _build_docx_from_template(content_str: str) -> bytes:
    """Generate a beautifully formatted, executive .docx Word document for passage templates."""
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    import io

    doc = Document()

    # Executive Page Setup: 0.8 in margins
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    # Global font styling
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    # Document Header Title
    tp = doc.add_paragraph()
    tp.paragraph_format.space_before = Pt(0)
    tp.paragraph_format.space_after = Pt(2)
    tr = tp.add_run("HearMeRead — Passage Assessment Template")
    tr.bold = True
    tr.font.size = Pt(16)
    tr.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    lines = content_str.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        trimmed = line.strip()

        # Notice Block Callout Card
        if "IMPORTANT NOTICE" in trimmed or (trimmed.startswith("=====") and i + 1 < len(lines) and "IMPORTANT NOTICE" in lines[i+1]):
            notice_lines = []
            if trimmed.startswith("====="):
                i += 1
            while i < len(lines) and not lines[i].strip().startswith("====="):
                t = lines[i].strip()
                if t:
                    notice_lines.append(t)
                i += 1
            if i < len(lines) and lines[i].strip().startswith("====="):
                i += 1

            tbl = doc.add_table(rows=1, cols=1)
            tbl.autofit = False
            tbl.columns[0].width = Inches(6.8)
            cell = tbl.cell(0, 0)
            _set_docx_cell_bg(cell, "FEF3C7")

            cp = cell.paragraphs[0]
            cp.paragraph_format.space_before = Pt(4)
            cp.paragraph_format.space_after = Pt(4)
            cr = cp.add_run("IMPORTANT NOTICE / PAUNAWA")
            cr.bold = True
            cr.font.size = Pt(11)
            cr.font.color.rgb = RGBColor(0x92, 0x40, 0x0E)

            for nl in notice_lines:
                if "IMPORTANT NOTICE" in nl:
                    continue
                cp2 = cell.add_paragraph()
                cp2.paragraph_format.space_before = Pt(2)
                cp2.paragraph_format.space_after = Pt(2)
                cr2 = cp2.add_run(nl)
                cr2.font.size = Pt(10)
                cr2.font.color.rgb = RGBColor(0x78, 0x35, 0x0F)

            sp = doc.add_paragraph()
            sp.paragraph_format.space_after = Pt(8)
            continue

        # Skip raw text separator lines
        if (trimmed.startswith("=====") or trimmed.startswith("-----")) and not ("SAMPLE FORMAT" in trimmed or "FILL OUT YOUR CONTENT" in trimmed):
            i += 1
            continue

        # Section Headings
        if "SAMPLE FORMAT" in trimmed or "MGA HALIMBAWA" in trimmed:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run("--------------------SAMPLE FORMAT / MGA HALIMBAWA (DO NOT EDIT SAMPLE)--------------------")
            r.bold = True
            r.font.size = Pt(11)
            tr.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
            i += 1
            continue

        if "FILL OUT YOUR CONTENT BELOW" in trimmed or "ISULAT ANG NILALAMAN DITO" in trimmed:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(20)
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run("==========FILL OUT YOUR CONTENT BELOW (ISULAT ANG NILALAMAN DITO)==========")
            r.bold = True
            r.font.size = Pt(12)
            tr.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
            i += 1
            continue

        if trimmed.startswith("---") and trimmed.endswith("---"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(trimmed)
            r.bold = True
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
            i += 1
            continue

        # Key labels
        labels = [
            "Language:", "Grade:", "Task 1:", "Task 2:", "Task 2 Words:",
            "Task 2 Sentences:", "Story Number:", "Title:", "Content:", "Questions:"
        ]
        is_key_label = any(trimmed.startswith(lbl) for lbl in labels)

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)

        if is_key_label:
            parts = trimmed.split(":", 1)
            r_lbl = p.add_run(parts[0] + ":")
            r_lbl.bold = True
            r_lbl.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
            if len(parts) > 1 and parts[1].strip():
                r_val = p.add_run(" " + parts[1].strip())
                r_val.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        elif trimmed in ["W:", "R:", "Q:", "A:"]:
            r_key = p.add_run(trimmed)
            r_key.bold = True
            r_key.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)
        else:
            r_txt = p.add_run(line)
            r_txt.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

        i += 1

    out = io.BytesIO()
    doc.save(out)
    out.seek(0)
    return out.getvalue()


# ── NEW: Download combined grade template (.txt or .docx) ──────────────────────
@router.get(
    "/template/a1",
    summary="Download a combined passage template (.txt or .docx) for a specific grade and language",
    response_class=Response,
)
async def download_a1_template(
    grade:       str = Query(..., description="grade_1 | grade_2 | grade_3"),
    language:    str = Query(..., description="filipino | english"),
    file_format: str = Query("txt", description="txt | docx"),
):
    key = (grade.lower().strip(), language.lower().strip())
    content = _COMBINED_TEMPLATES.get(key)
    if content is None:
        raise HTTPException(
            status_code=400,
            detail=(
                f"No template for grade='{grade}' language='{language}'. "
                "Valid combinations: grade_1/filipino, grade_2/filipino, "
                "grade_3/filipino, grade_3/english."
            ),
        )

    fmt = (file_format or "txt").lower().strip()
    if fmt == "docx":
        docx_bytes = _build_docx_from_template(content)
        filename = f"hearmeread_template_{grade}_{language}.docx"
        return Response(
            content=docx_bytes,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f"attachment; filename={filename}"},
        )

    filename = f"hearmeread_template_{grade}_{language}.txt"
    return Response(
        content=content.encode("utf-8"),
        media_type="text/plain; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


# ── Download A2 template route alias (backward compatible) ─────────────────────
@router.get(
    "/template/a2",
    summary="Download the Assessment template (.txt or .docx)",
    response_class=Response,
)
async def download_a2_template(
    file_format: str = Query("txt", description="txt | docx"),
):
    content = _COMBINED_GRADE2_FILIPINO
    fmt = (file_format or "txt").lower().strip()
    if fmt == "docx":
        docx_bytes = _build_docx_from_template(content)
        filename = "hearmeread_template_grade_2_filipino.docx"
        return Response(
            content=docx_bytes,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f"attachment; filename={filename}"},
        )

    return Response(
        content=content.encode("utf-8"),
        media_type="text/plain; charset=utf-8",
        headers={"Content-Disposition": "attachment; filename=hearmeread_template_grade_2_filipino.txt"},
    )


# ── Upload: combined passage + questions ──────────────────────────────────────
class CombinedUploadResponse(PassageResponse):
    imported_questions: List[QuestionResponse] = []


@router.post(
    "/upload",
    response_model=CombinedUploadResponse,
    status_code=status.HTTP_201_CREATED,
    summary="Upload a .docx or .txt file containing both passage and questions",
)
async def upload_combined(
    file:            UploadFile  = File(...),
    title:           str         = Form(..., min_length=1, max_length=255),
    language:        Language    = Form(...),
    grade_level:     GradeLevel  = Form(...),
    db:              AsyncSession = Depends(get_db),
    current_teacher: Teacher      = Depends(get_current_teacher),
):
    file_bytes = await file.read()
    validate_upload(file_bytes, file.filename)
    parsed = parse_combined(file_bytes, file.filename)
    passage = await passage_service.create_passage_from_docx(
        db=db, title=title, language=language, grade_level=grade_level,
        content=parsed.passage_content, teacher_id=current_teacher.id,
    )
    questions = await question_service.bulk_create_questions(
        db=db, passage_id=passage.id, texts=parsed.questions, teacher_id=current_teacher.id,
    )
    return CombinedUploadResponse(
        **PassageResponse.model_validate(passage).model_dump(),
        imported_questions=questions,
    )


# ── Upload: passage text only ─────────────────────────────────────────────────
@router.post(
    "/upload/passage-only",
    response_model=PassageResponse,
    status_code=status.HTTP_201_CREATED,
    summary="Upload a .docx or .txt file containing only the passage text",
)
async def upload_passage_only(
    file:            UploadFile  = File(...),
    title:           str         = Form(..., min_length=1, max_length=255),
    language:        Language    = Form(...),
    grade_level:     GradeLevel  = Form(...),
    db:              AsyncSession = Depends(get_db),
    current_teacher: Teacher      = Depends(get_current_teacher),
):
    file_bytes = await file.read()
    validate_upload(file_bytes, file.filename)
    content = parse_passage_only(file_bytes, file.filename)
    return await passage_service.create_passage_from_docx(
        db=db, title=title, language=language, grade_level=grade_level,
        content=content, teacher_id=current_teacher.id,
    )


# ── Bulk create passages (JSON) ───────────────────────────────────────────────
class BulkPassageQuestion(BaseModel):
    text:       str
    answer_key: Optional[str] = None
    order:      int           = 0


class BulkPassageItem(BaseModel):
    language:        Language
    grade_level:     Optional[GradeLevel] = None
    assessment_type: Optional[int]        = None
    title:           Optional[str]        = None
    content:         Optional[str]        = None
    task1_content:   Optional[str]        = None
    task2_words:     Optional[str]        = None
    task2_sentences: Optional[str]        = None
    story_number:    Optional[int]        = None
    questions:       List[BulkPassageQuestion] = []


class BulkCreateResult(BaseModel):
    index:      int
    passage_id: Optional[int] = None
    title:      Optional[str] = None
    error:      Optional[str] = None


class BulkCreateResponse(BaseModel):
    created: int
    failed:  int
    results: List[BulkCreateResult]


@router.post(
    "/bulk",
    response_model=BulkCreateResponse,
    status_code=status.HTTP_201_CREATED,
    summary="Create multiple passages with questions in a single request",
)
async def bulk_create_passages(
    items:           List[BulkPassageItem],
    db:              AsyncSession = Depends(get_db),
    current_teacher: Teacher      = Depends(get_current_teacher),
):
    """Create up to 20 passages (with inline questions) in one batched DB transaction."""
    if len(items) > 20:
        raise HTTPException(status_code=400, detail="Maximum 20 passages per bulk request.")

    from app.models import Passage as PassageModel, Question, PassageVisibility, ActivityLog

    created = 0
    failed  = 0
    results: List[BulkCreateResult] = []
    passages_to_add = []
    items_map = []  # (index, item, passage_obj) for question creation after flush

    # Pre-resolve titles for A1 items before attaching anything to session (avoids autoflush on dirty session)
    resolved_titles = {}
    for idx, item in enumerate(items):
        if item.assessment_type == 1 and not item.title and item.grade_level and item.language:
            lang_str = str(item.language.value if hasattr(item.language, "value") else item.language)
            grade_str = str(item.grade_level.value if hasattr(item.grade_level, "value") else item.grade_level)
            try:
                resolved_titles[idx] = await passage_service._a1_title(
                    db=db,
                    teacher_id=current_teacher.id,
                    language=lang_str,
                    grade_level=grade_str,
                )
            except Exception:
                resolved_titles[idx] = passage_service._a1_base_title(lang_str, grade_str)

    # Build passage objects
    for idx, item in enumerate(items):
        try:
            # Skip empty A1 items
            if item.assessment_type == 1:
                t1 = (item.task1_content or "").strip()
                t2w = (item.task2_words or "").strip()
                t2s = (item.task2_sentences or "").strip()
                # Clean out boilerplate if present
                if t2w.lower() in ("words:", "words:\n---", "words:\n\n---"):
                    t2w = ""
                if not t1 and not t2w and not t2s:
                    results.append(BulkCreateResult(index=idx, error="Empty Assessment 1 passage ignored."))
                    failed += 1
                    continue

            # Skip empty A2 items
            if item.assessment_type == 2:
                content = (item.content or "").strip()
                title = (item.title or "").strip()
                if not content and not title and not item.questions:
                    results.append(BulkCreateResult(index=idx, error="Empty Assessment 2 passage ignored."))
                    failed += 1
                    continue

            title = item.title or resolved_titles.get(idx)
            content = item.content or ""
            passage = PassageModel(
                teacher_id=current_teacher.id,
                title=title,
                content=content if content.strip() else None,
                language=item.language,
                grade_level=item.grade_level,
                word_count=len(content.split()) if content.strip() else 0,
                visibility=PassageVisibility.private,
                assessment_type=item.assessment_type,
                task1_content=item.task1_content,
                task2_words=item.task2_words,
                task2_sentences=item.task2_sentences,
                story_number=item.story_number,
            )
            db.add(passage)
            passages_to_add.append(passage)
            items_map.append((idx, item, passage))
        except Exception as exc:
            results.append(BulkCreateResult(index=idx, error=str(exc)))
            failed += 1

    if not passages_to_add:
        return BulkCreateResponse(created=0, failed=failed, results=results)

    try:
        # Flush to get auto-generated IDs for all passages
        await db.flush()

        # Create questions and activity logs in bulk
        all_questions = []
        all_logs = []
        for idx, item, passage in items_map:
            try:
                for q_idx, q in enumerate(item.questions):
                    if q.text.strip():
                        all_questions.append(Question(
                            passage_id=passage.id,
                            text=q.text.strip(),
                            answer_key=q.answer_key,
                            order=q.order if q.order else q_idx,
                        ))

                # Prepare activity log
                if current_teacher.school_id:
                    all_logs.append(ActivityLog(
                        teacher_id=current_teacher.id,
                        school_id=current_teacher.school_id,
                        action="uploaded_passage",
                        entity_type="passage",
                        entity_id=passage.id,
                        log_metadata={"title": passage.title or "Untitled"},
                    ))

                results.append(BulkCreateResult(
                    index=idx, passage_id=passage.id, title=passage.title,
                ))
                created += 1
            except Exception as exc:
                results.append(BulkCreateResult(index=idx, error=str(exc)))
                failed += 1

        if all_questions:
            db.add_all(all_questions)
        if all_logs:
            db.add_all(all_logs)

        # Single commit for everything
        await db.commit()
    except Exception as exc:
        await db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Database transaction error during bulk create: {str(exc)}",
        )

    return BulkCreateResponse(created=created, failed=failed, results=results)


# ── Upload original file to R2 ────────────────────────────────────────────────
@router.post(
    "/{passage_id}/file",
    status_code=status.HTTP_200_OK,
    summary="Upload the original passage file to R2 storage",
)
async def upload_passage_file(
    passage_id:      int,
    file:            UploadFile  = File(...),
    db:              AsyncSession = Depends(get_db),
    current_teacher: Teacher      = Depends(get_current_teacher),
):
    passage = await passage_service.get_passage_by_id(db, passage_id, current_teacher.id)
    file_bytes = await file.read()
    r2_key = await storage_service.upload_passage_file(
        file_bytes, file.filename or "passage", str(current_teacher.id)
    )
    passage.file_path = r2_key
    await db.commit()
    await db.refresh(passage)
    return {"file_path": r2_key}


# ── Get one ───────────────────────────────────────────────────────────────────
@router.get("/{passage_id}", response_model=PassageResponse, summary="Get a passage")
async def get_passage(
    passage_id:      int,
    db:              AsyncSession = Depends(get_db),
    current_teacher: Teacher      = Depends(get_current_teacher),
):
    return await passage_service.get_passage_by_id(
        db=db, passage_id=passage_id, teacher_id=current_teacher.id
    )


# ── Update ────────────────────────────────────────────────────────────────────
@router.patch("/{passage_id}", response_model=PassageResponse, summary="Update a passage")
async def update_passage(
    passage_id:      int,
    data:            PassageUpdate,
    db:              AsyncSession  = Depends(get_db),
    current_teacher: Teacher       = Depends(get_current_teacher),
):
    return await passage_service.update_passage(
        db=db, passage_id=passage_id, data=data, teacher_id=current_teacher.id
    )


# ── Archive ───────────────────────────────────────────────────────────────────
@router.delete(
    "/{passage_id}",
    status_code=status.HTTP_204_NO_CONTENT,
    summary="Archive a passage (soft delete)",
)
async def archive_passage(
    passage_id:      int,
    db:              AsyncSession = Depends(get_db),
    current_teacher: Teacher      = Depends(get_current_teacher),
):
    await passage_service.archive_passage(
        db=db, passage_id=passage_id, teacher_id=current_teacher.id
    )
